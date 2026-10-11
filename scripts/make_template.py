"""schema/sheet-schema.json 으로 구글 시트용 엑셀 템플릿 생성. 사용: python3 scripts/make_template.py"""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
root = pathlib.Path(__file__).resolve().parent.parent
schema = json.load(open(root / 'schema/sheet-schema.json', encoding='utf-8'))
wb = Workbook(); wb.remove(wb.active)
head = PatternFill('solid', fgColor='DDE3F5')
for tab, spec in schema.items():
    ws = wb.create_sheet(tab)
    for i, (key, _) in enumerate(spec['columns'], 1):
        c = ws.cell(row=1, column=i, value=key); c.font = Font(bold=True); c.fill = head
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(key) + 4)
    ws.freeze_panes = 'A2'
ex = wb.create_sheet('설명', 0)
ex.append(['탭 이름', '한글 이름', '열 이름', '설명'])
for c in ex[1]: c.font = Font(bold=True); c.fill = head
for tab, spec in schema.items():
    for key, desc in spec['columns']: ex.append([tab, spec['label'], key, desc])
for col, w in zip('ABCD', (18, 16, 20, 70)): ex.column_dimensions[col].width = w
ex.freeze_panes = 'A2'
ws = wb['config']
for row in (['brand_KK', '키드키즈', ''], ['brand_SE', '실버에듀넷', ''], ['channel_IG', '인스타그램', ''], ['channel_YT', '유튜브', ''], ['channel_BL', '네이버 블로그', '']): ws.append(row)
out = root / 'sheet-template/SNS대시보드_시트양식.xlsx'; wb.save(out); print('saved', out)
